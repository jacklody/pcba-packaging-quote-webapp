from packaging_core.api.schema import BusinessInput, ConstraintsInput, LogisticsInput, PackagingRequest, ProductInput
from packaging_core.api.service import estimate_packaging
from packaging_core.export.csv_exporter import result_to_csv
from packaging_core.export.excel_exporter import result_to_workbook


def make_request():
    return PackagingRequest(
        ProductInput("DRV-26P", "assembly", 620, 360, 95, 9.8, "medium", True),
        BusinessInput(5000, True, True),
        LogisticsInput("pallet", "深圳市", "特惠普运", "1100x1100", 1500, 700),
        ConstraintsInput([1, 2, 4, 6], 25),
    )


def test_csv_contains_bom_rows_and_total():
    result = estimate_packaging(make_request())

    csv_text = result_to_csv(result)

    assert "项目,规格,数量,单价(元),成本(元/台),置信度" in csv_text
    assert "EPS左缓冲" in csv_text
    assert "5层加厚纸箱" in csv_text
    assert "合计" in csv_text
    # utf-8 BOM already stripped by str; ensure no surrogate issues on round-trip
    csv_text.encode("utf-8")


def test_excel_workbook_has_sheets_and_values(tmp_path):
    result = estimate_packaging(make_request())
    target = tmp_path / "packaging_result.xlsx"

    result_to_workbook(result, str(target))

    from openpyxl import load_workbook

    wb = load_workbook(str(target))
    assert wb.sheetnames == ["方案概览", "包材BOM", "风险与备选"]

    bom_ws = wb["包材BOM"]
    bom_items = {row[0] for row in bom_ws.iter_rows(min_row=2, values_only=True) if row[0]}
    assert "EPS左缓冲" in bom_items
    assert "5层加厚纸箱" in bom_items

    overview_ws = wb["方案概览"]
    text_values = [str(value) for row in overview_ws.iter_rows(values_only=True) for value in row if value is not None]
    assert any("EPS成型缓冲" in v for v in text_values)

    risk_ws = wb["风险与备选"]
    risk_text = [str(value) for row in risk_ws.iter_rows(values_only=True) for value in row if value is not None]
    assert any("风险" in v or "EPE" in v for v in risk_text)
