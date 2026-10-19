from openpyxl import Workbook


def _write_bom_sheet(ws, bom):
    headers = ["项目", "规格", "数量", "单价(元)", "成本(元/台)", "置信度", "公式"]
    ws.append(headers)
    for row in bom:
        ws.append([row["item"], row["spec"], row["qty"], row["unit_price"], row["cost_per_unit"], row["confidence"], row.get("formula", "")])
    ws.append(["合计", "", "", "", round(sum(r["cost_per_unit"] for r in bom), 2), "", ""])


def _write_overview_sheet(ws, result):
    recommended = result["recommended"]
    cost = result["cost"]
    ws.append(["推荐方案", recommended["summary"]])
    ws.append(["台/箱", recommended["units_per_box"]])
    ws.append(["箱/托", recommended["boxes_per_pallet"]])
    ws.append(["台/托", recommended["units_per_pallet"]])
    ws.append(["总成本(元/台)", recommended["total_cost_per_unit"]])
    ws.append(["置信度", recommended["confidence"]])
    ws.append([])
    ws.append(["成本拆解", "元/台"])
    ws.append(["包材小计", cost["packaging_material_per_unit"]])
    ws.append(["模具摊销", cost["mold_amortize_per_unit"]])
    ws.append(["托盘辅料", cost["pallet_accessory_per_unit"]])
    ws.append(["运费", cost["freight_per_unit"]])


def _write_risk_sheet(ws, risks, alternatives):
    ws.append(["风险提示"])
    ws.append(["级别", "类型", "说明"])
    for risk in risks:
        ws.append([risk["level"], risk["type"], risk["message"]])
    ws.append([])
    ws.append(["备选方案"])
    ws.append(["方案", "说明", "元/台", "置信度"])
    for alt in alternatives:
        ws.append([alt["scheme_id"], alt["summary"], alt["total_cost_per_unit"], alt["confidence"]])


def result_to_workbook(result: dict, path: str) -> None:
    wb = Workbook()
    overview_ws = wb.active
    overview_ws.title = "方案概览"
    _write_overview_sheet(overview_ws, result)
    _write_bom_sheet(wb.create_sheet("包材BOM"), result["bom"])
    _write_risk_sheet(wb.create_sheet("风险与备选"), result["risks"], result["alternatives"])
    wb.save(path)
